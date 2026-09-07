from docx import Document

from axp_core.rich_document import DOCX_RICH_EXTRACTOR_VERSION, extract_docx


def test_physical_vertical_merge_content_is_emitted_once(tmp_path):
    path = tmp_path / "merged.docx"
    document = Document()
    table = document.add_table(rows=4, cols=2)
    table.cell(0, 0).text = "MONOGRAPHIE ANALYTIQUE"
    table.cell(0, 0).merge(table.cell(3, 0))
    document.save(path)

    rich, _ = extract_docx(path, document_id=1, filename=path.name, sha256="abc")
    manifest = rich.to_manifest()
    cells = [row["cells"][0] for row in manifest["blocks"][0]["rows"]]
    assert [cell["v_merge"] for cell in cells] == ["restart", "continue", "continue", "continue"]
    assert str(manifest["blocks"]).count("MONOGRAPHIE ANALYTIQUE") == 1
    assert manifest["extractor_version"] == DOCX_RICH_EXTRACTOR_VERSION == 2
    assert manifest["document"]["fidelity"] == "full"
    assert manifest["diagnostics"]["unsupported_by_type"] == {}
