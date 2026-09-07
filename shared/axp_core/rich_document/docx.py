from __future__ import annotations

import hashlib
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from .model import RICH_DOCUMENT_SCHEMA_VERSION, RichAsset, RichDocument
from .safety import (MAX_BLOCKS, MAX_NESTING_DEPTH, MAX_TABLE_CELLS, MAX_TEXT_CHARACTERS,
                     RichDocumentError, validate_docx)

MEDIA = {"png": ("image/png", True), "jpg": ("image/jpeg", True), "jpeg": ("image/jpeg", True),
         "gif": ("image/gif", True), "bmp": ("image/bmp", True), "webp": ("image/webp", True),
         "tif": ("image/tiff", False), "tiff": ("image/tiff", False), "svg": ("image/svg+xml", True),
         "emf": ("image/x-emf", False), "wmf": ("image/x-wmf", False)}
UNSUPPORTED_TYPES = {"chart": "chart", "diagram": "drawing", "oleObject": "ole",
                     "object": "ole", "txbxContent": "text_box", "control": "control",
                     "oMath": "equation", "oMathPara": "equation"}


class _Extractor:
    def __init__(self, document, document_id, filename, source_sha):
        self.document, self.document_id, self.filename, self.source_sha = document, document_id, filename, source_sha
        self.assets, self.blobs, self.blocks = {}, {}, []
        self.text_chars = self.table_cells = self.unsupported = self.external = 0
        self.unsupported_by_type = {}
        self._unsupported_nodes = set()

    def unsupported_block(self, node, object_type=None):
        if node in self._unsupported_nodes:
            return None
        self._unsupported_nodes.add(node)
        kind = object_type or UNSUPPORTED_TYPES.get(node.tag.rsplit("}", 1)[-1], "unknown")
        self.unsupported += 1
        self.unsupported_by_type[kind] = self.unsupported_by_type.get(kind, 0) + 1
        return {"type": "unsupported", "object_type": kind}

    def limit(self):
        if len(self.blocks) > MAX_BLOCKS or self.text_chars > MAX_TEXT_CHARACTERS or self.table_cells > MAX_TABLE_CELLS:
            raise RichDocumentError("rich_document_too_complex")

    def image(self, blip, drawing):
        rid = blip.get(qn("r:embed"))
        link = blip.get(qn("r:link"))
        if not rid:
            if link:
                self.external += 1
            else:
                self.unsupported += 1
            return None
        try:
            part = self.document.part.related_parts[rid]
            blob = part.blob
        except (KeyError, AttributeError):
            self.unsupported += 1
            return None
        digest = hashlib.sha256(blob).hexdigest()
        asset_id = "img-" + digest[:32]
        ext = Path(str(getattr(part, "partname", ""))).suffix.lower().lstrip(".")
        media_type, browser = MEDIA.get(ext, (getattr(part, "content_type", None) or "application/octet-stream", False))
        extent = drawing.xpath("ancestor-or-self::*[local-name()='inline' or local-name()='anchor']/*[local-name()='extent']")
        props = drawing.xpath("ancestor-or-self::*[local-name()='inline' or local-name()='anchor']/*[local-name()='docPr']")
        width = int(extent[0].get("cx")) if extent and extent[0].get("cx") else None
        height = int(extent[0].get("cy")) if extent and extent[0].get("cy") else None
        alt = props[0].get("descr") if props else None
        title = props[0].get("title") if props else None
        if asset_id not in self.assets:
            suffix = {"image/jpeg": ".jpeg", "image/tiff": ".tiff", "image/svg+xml": ".svg",
                      "image/x-emf": ".emf", "image/x-wmf": ".wmf"}.get(media_type, "." + ext if ext else ".bin")
            self.assets[asset_id] = RichAsset(asset_id, media_type, len(blob), digest, rid, alt, title,
                                              width, height, browser, asset_id + suffix)
            self.blobs[asset_id] = blob
        return {"type": "image", "asset_id": asset_id, "relationship_id": rid,
                "alt_text": alt, "title": title, "width_emu": width, "height_emu": height}

    def paragraph_content(self, paragraph):
        content = []
        for child in paragraph._p:
            nodes = list(child.iter()) if child.tag in (qn("w:r"), qn("w:hyperlink")) else []
            hyperlink = child if child.tag == qn("w:hyperlink") else None
            href = None
            if hyperlink is not None:
                rid = hyperlink.get(qn("r:id"))
                try:
                    rel = paragraph.part.rels[rid] if rid else None
                    href = rel.target_ref if rel is not None else None
                    if rel is not None and rel.is_external:
                        self.external += 1
                except KeyError:
                    pass
            for node in nodes:
                if node.tag == qn("w:t"):
                    run = next((x for x in node.iterancestors() if x.tag == qn("w:r")), None)
                    props = run.find(qn("w:rPr")) if run is not None else None
                    item = {"type": "text", "text": node.text or "", "bold": props is not None and props.find(qn("w:b")) is not None,
                            "italic": props is not None and props.find(qn("w:i")) is not None,
                            "underline": props is not None and props.find(qn("w:u")) is not None}
                    if href: item["hyperlink"] = href
                    content.append(item); self.text_chars += len(item["text"])
                elif node.tag == qn("w:tab"):
                    content.append({"type": "text", "text": "\t", "bold": False, "italic": False, "underline": False})
                    self.text_chars += 1
                elif node.tag == qn("w:br") and node.get(qn("w:type")) == "page":
                    content.append({"type": "page_break"})
                elif node.tag == qn("a:blip"):
                    item = self.image(node, node)
                    if item: content.append(item)
                else:
                    local = node.tag.rsplit("}", 1)[-1]
                    if local in UNSUPPORTED_TYPES:
                        item = self.unsupported_block(node)
                        if item: content.append(item)
            # A DrawingML container with neither an image nor a more specific
            # supported/unsupported child must not disappear silently.
            if child.tag == qn("w:r"):
                for drawing in child.xpath(".//*[local-name()='drawing']"):
                    if not drawing.xpath(".//*[local-name()='blip' or local-name()='chart' or local-name()='diagram']"):
                        item = self.unsupported_block(drawing, "drawing")
                        if item: content.append(item)
        return content

    def paragraph(self, element, parent, depth=0):
        paragraph = Paragraph(element, parent)
        content = self.paragraph_content(paragraph)
        style = paragraph.style.name if paragraph.style is not None else None
        match = re.fullmatch(r"Heading\s+([1-9])", style or "", re.IGNORECASE)
        if match and all(item["type"] == "text" for item in content):
            return [{"type": "heading", "level": int(match.group(1)), "text": "".join(x["text"] for x in content), "style": style}]
        # A standalone drawing is a body image; mixed drawings remain ordered inline content.
        meaningful_text = any(x["type"] == "text" and x["text"] for x in content)
        if not meaningful_text and len(content) == 1 and content[0]["type"] == "image":
            return [content[0]]
        if any(item["type"] == "page_break" for item in content):
            blocks, segment = [], []
            for item in content:
                if item["type"] == "page_break":
                    if segment: blocks.append({"type": "paragraph", "style": style, "content": segment})
                    blocks.append({"type": "page_break"}); segment = []
                else:
                    segment.append(item)
            if segment: blocks.append({"type": "paragraph", "style": style, "content": segment})
            return blocks
        return [{"type": "paragraph", "style": style, "content": content}]

    def table(self, element, parent, depth):
        if depth > MAX_NESTING_DEPTH: raise RichDocumentError("rich_document_too_complex")
        table = Table(element, parent)
        rows = []
        for row in table._tbl.tr_lst:
            cells = []
            # tc_lst contains physical XML cells.  ``row.cells`` expands grid
            # spans and vertical merges into repeated logical proxies.
            for tc in row.tc_lst:
                cell = _Cell(tc, table)
                self.table_cells += 1
                tc_pr = tc.tcPr
                span = tc_pr.gridSpan.val if tc_pr is not None and tc_pr.gridSpan is not None else 1
                merge_element = tc_pr.vMerge if tc_pr is not None else None
                merge = ("continue" if merge_element is not None and merge_element.val in (None, "continue")
                         else "restart" if merge_element is not None and merge_element.val == "restart" else "none")
                blocks = self.walk(tc, cell, depth + 1)
                cells.append({"blocks": blocks, "grid_span": int(span), "v_merge": merge})
            rows.append({"cells": cells})
        style = table.style.name if table.style is not None else None
        return {"type": "table", "style": style, "rows": rows}

    def walk(self, container, parent, depth=0):
        result = []
        for child in container.iterchildren():
            if child.tag == qn("w:p"): result.extend(self.paragraph(child, parent, depth))
            elif child.tag == qn("w:tbl"): result.append(self.table(child, parent, depth))
        return result

    def run(self):
        self.blocks = self.walk(self.document.element.body, self.document)
        xml = self.document.element.body
        # Count objects outside paragraphs too, without double counting located placeholders.
        for node in xml.iter():
            if node.tag.rsplit("}", 1)[-1] in UNSUPPORTED_TYPES:
                self.unsupported_block(node)
        self.limit()
        counts = {"paragraphs": sum(x["type"] == "paragraph" for x in self.blocks),
                  "tables": sum(x["type"] == "table" for x in self.blocks),
                  "images": len(self.assets), "image_references": self._image_references(self.blocks),
                  "unsupported_objects": self.unsupported,
                  "unsupported_by_type": dict(sorted(self.unsupported_by_type.items())),
                  "external_relationships": self.external, "truncation": False,
                  "parser_warning_count": 0, "duplicates_suppressed": 0,
                  "block_count": len(self.blocks)}
        fidelity = "full" if (not self.unsupported and not counts["truncation"]
                              and counts["parser_warning_count"] == 0) else "partial"
        return RichDocument(RICH_DOCUMENT_SCHEMA_VERSION, self.document_id, self.filename, self.source_sha,
                            tuple(self.blocks), tuple(self.assets.values()), fidelity, counts), self.blobs

    def _image_references(self, values):
        total = 0
        for value in values:
            if isinstance(value, dict):
                total += value.get("type") == "image"
                total += self._image_references(value.values())
            elif isinstance(value, (list, tuple)): total += self._image_references(value)
        return total


def extract_docx(path, *, document_id, filename, sha256):
    source = Path(path)
    validate_docx(source)
    try:
        document = Document(source)
        return _Extractor(document, document_id, filename, sha256).run()
    except RichDocumentError:
        raise
    except Exception as exc:
        raise RichDocumentError("rich_document_invalid") from exc
